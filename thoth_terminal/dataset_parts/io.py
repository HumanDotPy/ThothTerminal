from __future__ import annotations

import csv
import hashlib
import json
import re
from datetime import date, datetime
from io import BytesIO, StringIO
from threading import RLock
from typing import Any

from sqlalchemy import func
from sqlalchemy.orm import Session

from thoth_terminal.db import utcnow
from thoth_terminal.dataset_catalog import (
    DATASET_FAMILY_COLUMN_TEMPLATES,
    DATASET_FAMILY_MATERIALIZER_FAMILIES,
    DATASET_FAMILY_META,
    DATASET_FAMILY_REQUIRED_ROLES,
    DATASET_OBJECT_TYPES,
    DATASET_VIEW_OPERATOR_VALUES,
    DATASET_VIEW_OPERATORS,
)
from thoth_terminal.models import (
    ChartDatasetView,
    Dataset,
    DatasetColumn,
    Document,
    DatasetLink,
    DatasetRow,
    DatasetVersion,
    DatasetView,
    InstrumentRequest,
    ProductStageEntry,
    StructuredDataFact,
)
from thoth_terminal.product_stage_service import serialize_product_stage_entry


def parse_csv_dataset(text: str) -> tuple[list[dict[str, Any]], list[dict[str, Any]]]:
    reader = csv.DictReader(StringIO(text))
    if not reader.fieldnames:
        return [], []
    rows = []
    for row in reader:
        rows.append({str(key or "").strip(): _coerce_scalar(value) for key, value in row.items()})
    return _normalize_columns_and_rows(None, rows)


def parse_xlsx_dataset(blob: bytes) -> tuple[list[dict[str, Any]], list[dict[str, Any]]]:
    try:
        from openpyxl import load_workbook
    except Exception as exc:  # pragma: no cover - dependency guarded at runtime
        raise RuntimeError("XLSX support requires openpyxl.") from exc
    wb = load_workbook(filename=BytesIO(blob), read_only=True, data_only=True)
    ws = wb.active
    rows_iter = list(ws.iter_rows(values_only=True))
    if not rows_iter:
        return [], []
    headers = [str(cell or "").strip() or f"column_{idx + 1}" for idx, cell in enumerate(rows_iter[0])]
    data_rows = []
    for raw_row in rows_iter[1:]:
        data_rows.append({headers[idx]: _coerce_scalar(val) for idx, val in enumerate(raw_row or [])})
    return _normalize_columns_and_rows(None, data_rows)


def export_dataset_csv_bytes(
    session: Session,
    dataset: Dataset,
    *,
    view_id: int | None = None,
    version_id: int | None = None,
) -> bytes:
    payload = query_dataset(session, dataset_id=int(dataset.id), version_id=version_id, view_id=view_id)
    version = session.get(DatasetVersion, int(payload["version"]["id"])) if payload.get("version") else None
    cols = session.query(DatasetColumn).filter_by(dataset_version_id=int(version.id)).order_by(DatasetColumn.sort_order.asc()).all() if version else []
    keys = [col.key for col in cols] if cols else list((payload["rows"][0]["values"].keys() if payload["rows"] else []))
    buf = StringIO()
    writer = csv.DictWriter(buf, fieldnames=keys)
    writer.writeheader()
    for row in payload["rows"]:
        writer.writerow({key: row["values"].get(key) for key in keys})
    return buf.getvalue().encode("utf-8")


def export_dataset_xlsx_bytes(
    session: Session,
    dataset: Dataset,
    *,
    view_id: int | None = None,
    version_id: int | None = None,
) -> bytes:
    try:
        from openpyxl import Workbook
    except Exception as exc:  # pragma: no cover - dependency guarded at runtime
        raise RuntimeError("XLSX support requires openpyxl.") from exc
    payload = query_dataset(session, dataset_id=int(dataset.id), version_id=version_id, view_id=view_id)
    version = session.get(DatasetVersion, int(payload["version"]["id"])) if payload.get("version") else None
    cols = session.query(DatasetColumn).filter_by(dataset_version_id=int(version.id)).order_by(DatasetColumn.sort_order.asc()).all() if version else []
    keys = [col.key for col in cols] if cols else list((payload["rows"][0]["values"].keys() if payload["rows"] else []))
    wb = Workbook()
    ws = wb.active
    ws.title = "Dataset"
    ws.append(keys)
    for row in payload["rows"]:
        ws.append([row["values"].get(key) for key in keys])
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()
